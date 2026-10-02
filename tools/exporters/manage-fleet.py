"""Fleet access audit. Credentials stay in memory; reports contain no credentials.

Private address/host-key reports are kept outside the repository. This command
does not change any remote service; installation is a separate explicit mode.
"""
import argparse, base64, concurrent.futures, hashlib, json, logging, os, pathlib, re, shlex, socket, threading, time
import openpyxl, paramiko
from urllib.parse import urlsplit

ROOT = pathlib.Path(__file__).resolve().parents[2]
PRIVATE = pathlib.Path(os.environ.get('TEMP', '/tmp')) / 'hpclab-exporter-private'
KEY = pathlib.Path.home() / '.ssh' / 'pk-lab-2026'
logging.getLogger('paramiko').setLevel(logging.CRITICAL)
HOST_LOCK = threading.Lock()

REMOTE_AUDIT = r'''
import json, os, platform, subprocess, urllib.request
def run(args, timeout=6):
 try:
  p=subprocess.run(args,capture_output=True,text=True,timeout=timeout)
  return p.returncode,p.stdout.strip()
 except Exception:return -1,''
release={}
try:
 for line in open('/etc/os-release'):
  if '=' in line:
   k,v=line.strip().split('=',1);release[k]=v.strip('"')
except Exception:pass
ports=[]
code,out=run(['ss','-lntH'])
for line in out.splitlines():
 fields=line.split()
 if len(fields)>3:
  try:ports.append(int(fields[3].rsplit(':',1)[1]))
  except Exception:pass
services=[]
code,out=run(['systemctl','list-units','--all','--type=service','--no-legend','--no-pager'])
for line in out.splitlines():
 if 'exporter' in line:
  fields=line.split()
  if fields:services.append(fields[0])
metrics={}
for port in [9100,9835,9400,9401]:
 try:
  body=urllib.request.build_opener(urllib.request.ProxyHandler({})).open('http://127.0.0.1:%d/metrics'%port,timeout=2).read(200000).decode('utf8','replace')
  metrics[str(port)]={'node':'node_cpu_seconds_total' in body,'nvidia':'nvidia_smi_gpu_info' in body,'dcgm':'DCGM_FI_DEV' in body}
 except Exception:pass
code,gpu=run(['nvidia-smi','--query-gpu=index,name,uuid,driver_version,mig.mode.current','--format=csv,noheader,nounits'])
gpuRows=[]
if code==0:
 for line in gpu.splitlines():
  fields=[x.strip() for x in line.split(',')]
  if len(fields)==5:gpuRows.append(dict(zip(['index','name','uuid','driver','mig'],fields)))
_,processes=run(['ps','-eo','comm='])
print(json.dumps({'os':release.get('ID'), 'version':release.get('VERSION_ID'), 'kernel':platform.system(), 'arch':platform.machine(), 'uid':os.geteuid(), 'systemd':os.path.isdir('/run/systemd/system'), 'ports':sorted(set(ports)), 'services':services, 'metrics':metrics, 'gpus':gpuRows, 'nvidiaQuery':code, 'profilingActive':any(x.strip() in ['nsys','ncu','nvprof'] for x in processes.splitlines()), 'dcgmInstalled':run(['sh','-c','command -v dcgmi'])[0]==0, 'tegrastatsInstalled':run(['sh','-c','command -v tegrastats'])[0]==0, 'dockerAvailable':run(['sh','-c','command -v docker'])[0]==0}))
'''

class HostPolicy(paramiko.MissingHostKeyPolicy):
    def missing_host_key(self, client, hostname, key):
        with HOST_LOCK:
            file=PRIVATE/'known_hosts'
            keys=paramiko.HostKeys(str(file)) if file.exists() else paramiko.HostKeys()
            known=keys.lookup(hostname)
            if known and key.get_name() in known and known[key.get_name()] != key:
                raise paramiko.BadHostKeyException(hostname,key,known[key.get_name()])
            keys.add(hostname,key.get_name(),key);keys.save(str(file))
            client.get_host_keys().add(hostname,key.get_name(),key)

def credentials(sheet, row):
    if not row:return None,None
    user=sheet.cell(row,14).value;password=sheet.cell(row,15).value
    if user is None:return None,None
    user=str(user).strip()
    if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.@-]{0,63}',user):return None,None
    if isinstance(password,float) and password.is_integer():password=int(password)
    return user,str(password) if password is not None else None

def load_targets(xlsx, include_hidden=False):
    inventory=json.loads((ROOT/'lab-inventory.default.json').read_text(encoding='utf8'))['machines']
    raw=json.loads((ROOT/'tools/data/lab-inventory.raw.json').read_text(encoding='utf8'))['machines']
    workbook=openpyxl.load_workbook(xlsx,read_only=True,data_only=True);sheet=workbook['Device List']
    targets=[]
    for i,m in enumerate(inventory):
        if not m.get('ip') or (not include_hidden and not m.get('include')):continue
        row=raw[i].get('row') if i<len(raw) else None
        user,password=credentials(sheet,row)
        targets.append({**m,'_user':user,'_password':password})
    workbook.close()
    return targets

def connect(target):
    ip=target['ip'];ports=[p for p in target.get('ports',[]) if p in (22,30678)] or [22,30678]
    reachable=[]
    for port in ports:
        try:
            with socket.create_connection((ip,port),timeout=1.5):reachable.append(port)
        except OSError:pass
    if not reachable:return None,None,None,'ssh-unreachable'
    users=[target['_user']] if target['_user'] else ['ubuntu','root']
    last='authentication-failed'
    for port in reachable:
        for user in users:
            client=paramiko.SSHClient()
            known=pathlib.Path.home()/'.ssh/known_hosts'
            if known.exists():client.load_system_host_keys(str(known))
            client.set_missing_host_key_policy(HostPolicy())
            try:
                client.connect(ip,port=port,username=user,password=target['_password'],key_filename=str(KEY) if KEY.exists() else None,look_for_keys=False,allow_agent=False,timeout=4,auth_timeout=5,banner_timeout=5)
                return client,user,port,None
            except paramiko.BadHostKeyException:
                client.close();return None,None,None,'host-key-changed'
            except paramiko.AuthenticationException:last='authentication-failed'
            except Exception:last='ssh-error'
            client.close()
    return None,None,None,last

def command(client, command, stdin=None, timeout=25):
    inp,out,err=client.exec_command(command,timeout=timeout)
    if stdin is not None:inp.write(stdin);inp.flush();inp.channel.shutdown_write()
    text=out.read().decode('utf8','replace');err.read()
    return out.channel.recv_exit_status(),text

def audit(target):
    public={'id':target['id'],'label':target['label'],'virtual':target.get('virtual',False)}
    client,user,port,error=connect(target)
    if error:return {**public,'status':error}
    try:
        payload=base64.b64encode(REMOTE_AUDIT.encode()).decode()
        code,text=command(client,"python3 -c \"import base64;exec(base64.b64decode('%s'))\""%payload)
        if code!=0:return {**public,'status':'python3-unavailable','sshPort':port}
        details=json.loads(text.strip().splitlines()[-1])
        sudo='root' if details['uid']==0 else None
        if not sudo:
            code,out=command(client,'sudo -n /usr/bin/id -u')
            if code==0 and out.strip()=='0':sudo='passwordless'
            elif target['_password'] is not None:
                code,out=command(client,"sudo -S -p '' /usr/bin/id -u",target['_password']+'\n')
                if code==0 and out.strip()=='0':sudo='password'
        return {**public,'status':'accessible','sshPort':port,'sudo':sudo,**details}
    except Exception:return {**public,'status':'audit-error','sshPort':port}
    finally:client.close()

def privileged(client, target, sudo, cmd, timeout=120):
    if sudo == 'root':return command(client,cmd,timeout=timeout)
    prefix="sudo -S -p '' " if sudo=='password' else 'sudo -n '
    return command(client,prefix+cmd,(target['_password']+'\n') if sudo=='password' else None,timeout=timeout)

def install(target, prior):
    result={'id':target['id'],'label':target['label'],'status':'installation-skipped'}
    if prior.get('status')!='accessible' or not prior.get('sudo'):
        return {**result,'reason':'administrative-access-unavailable'}
    if not prior.get('systemd') or prior.get('arch') not in ('x86_64','aarch64','arm64'):
        return {**result,'reason':'unsupported-platform'}
    client,user,port,error=connect(target)
    if error:return {**result,'reason':error}
    stage=None
    try:
        code,out=command(client,'mktemp -d /tmp/hpclab-exporter-stage.XXXXXXXX')
        stage=out.strip()
        if code or not re.fullmatch('/tmp/hpclab-exporter-stage\\.[A-Za-z0-9]{8}',stage):
            return {**result,'reason':'staging-unavailable'}
        manifest=json.loads((PRIVATE/'releases/manifest.json').read_text())
        arch='arm64' if prior['arch'] in ('aarch64','arm64') else 'x86_64'
        artifacts=[a for a in manifest['artifacts'] if a['arch']==arch]
        args=['bash',stage+'/install-linux.sh','--listen',target['ip']]
        sftp=client.open_sftp()
        try:
            for name in ['install-linux.sh','nvidia_smi_exporter.py','restrict-linux.sh']:
                # Normalize CRLF, ensuring scripts also run from Windows checkouts.
                with sftp.file(stage+'/'+name,'wb') as handle:
                    handle.write((ROOT/'tools/exporters'/name).read_bytes().replace(b'\r\n',b'\n'))
            for a in artifacts:
                contents=pathlib.Path(a['path']).read_bytes()
                if hashlib.sha256(contents).hexdigest()!=a['sha256']:
                    return {**result,'reason':'local-checksum-mismatch'}
                sftp.put(a['path'],stage+'/'+a['asset'])
                args.extend(['--'+a['kind']+'-version',a['version'],'--'+a['kind']+'-sha256',a['sha256'],'--'+a['kind']+'-archive',stage+'/'+a['asset']])
        finally:sftp.close()
        # The public-addressed fleet must not gain unrestricted exporter ports.
        # Limit only newly opened ports to the independent Prometheus and website nodes.
        ports=[str(p) for p in (9100,9835) if p not in prior.get('ports',[]) and (p==9100 or prior.get('gpus'))]
        if ports:
            nodes=json.loads((PRIVATE/'kube-nodes.json').read_text(encoding='utf-8-sig'))['items']
            sources={a['address'] for n in nodes for a in n['status'].get('addresses',[]) if a['type'] in ('InternalIP','ExternalIP')}
            sources.add(urlsplit((PRIVATE/'PROMETHEUS_URL.txt').read_text().strip()).hostname)
            if not all(re.fullmatch('[0-9.]+',s or '') for s in sources):
                return {**result,'reason':'monitoring-source-addresses-unavailable'}
            argsAcl=['bash',stage+'/restrict-linux.sh',target['ip'],','.join(ports),*sorted(sources)]
            code,out=privileged(client,target,prior['sudo'],shlex.join(argsAcl))
            if code:return {**result,'reason':'source-access-setup-failed'}
        code,out=privileged(client,target,prior['sudo'],shlex.join(args))
        messages=[line for line in out.splitlines() if line.startswith(('platform:','node:','gpu:','ERROR:'))]
        # Installation text is fixed and does not include network addresses or credentials.
        result.update(status='installer-completed' if code==0 else 'installer-error',messages=messages)
        check=r'''
import json,urllib.request,subprocess
host=HOST
result={}
for kind,port,metric,unit in [('node',9100,'node_cpu_seconds_total','hpclab-node-exporter'),('gpu',9835,'nvidia_smi_gpu_info','hpclab-nvidia-smi-exporter')]:
 try:
  body=urllib.request.build_opener(urllib.request.ProxyHandler({})).open('http://%s:%d/metrics'%(host,port),timeout=4).read(2000000).decode()
  lines=[x for x in body.splitlines() if x and not x.startswith('#')]
  result[kind]={'metricsValid':metric in body,'sampleCount':len(lines),'serviceActive':subprocess.run(['systemctl','is-active',unit],capture_output=True,text=True).stdout.strip()=='active'}
 except Exception:result[kind]={'metricsValid':False}
print(json.dumps(result))
'''.replace('HOST',repr(target['ip']))
        payload=base64.b64encode(check.encode()).decode()
        code,out=command(client,"python3 -c \"import base64;exec(base64.b64decode('%s'))\""%payload)
        if code==0:result['verification']=json.loads(out.strip().splitlines()[-1])
        return result
    except Exception:return {**result,'status':'installation-error','reason':'remote-operation-failed'}
    finally:
        if stage and re.fullmatch('/tmp/hpclab-exporter-stage\\.[A-Za-z0-9]{8}',stage):
            try:command(client,'rm -rf -- '+shlex.quote(stage))
            except Exception:pass
        client.close()

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--xlsx',default='C:/Users/PKH/Downloads/HPC Device List.xlsx');parser.add_argument('--include-hidden',action='store_true');parser.add_argument('--ids',nargs='*');parser.add_argument('--install',action='store_true');args=parser.parse_args()
    PRIVATE.mkdir(parents=True,exist_ok=True)
    targets=load_targets(args.xlsx,args.include_hidden)
    if args.ids:targets=[target for target in targets if target['id'] in args.ids]
    if args.install:
        if not args.ids:parser.error('--install requires explicit --ids of reviewed targets')
        prior={r['id']:r for r in json.loads((PRIVATE/'fleet-audit.json').read_text())}
        file=PRIVATE/'fleet-installs.json'
        installed={r['id']:r for r in json.loads(file.read_text())} if file.exists() else {}
        with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
            futures={pool.submit(install,target,prior.get(target['id'],{})):target for target in targets}
            for future in concurrent.futures.as_completed(futures):
                result=future.result();installed[result['id']]=result
                file.write_text(json.dumps(list(installed.values()),indent=2,ensure_ascii=True),encoding='utf8')
                print(json.dumps(result,ensure_ascii=True),flush=True)
        return
    results=[]
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
        futures={pool.submit(audit,target):target for target in targets}
        for future in concurrent.futures.as_completed(futures):
            result=future.result();results.append(result)
            print(json.dumps({k:result.get(k) for k in ['id','status','os','arch','sudo','metrics','gpus']},ensure_ascii=True),flush=True)
    results.sort(key=lambda x:x['id'])
    (PRIVATE/'fleet-audit.json').write_text(json.dumps(results,indent=2,ensure_ascii=True),encoding='utf8')
    print(json.dumps({'audited':len(results),'accessible':sum(r['status']=='accessible' for r in results),'administrable':sum(r.get('sudo') is not None for r in results)}),flush=True)

if __name__=='__main__':main()
